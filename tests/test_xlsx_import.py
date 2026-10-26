import pytest
import openpyxl
from sqlalchemy import delete, select

from routefit import db, xlsx_import, xlsx_template
from routefit.data import invalidate

GOOD = {
    # 실제 가게 데이터(data/places.json 의 s_001~)와 겹치지 않게 테스트 가게는 s_9xx 를 쓴다
    "id": "s_901", "name": "테스트 국밥", "category": "식당", "area": "동구 동명동", "address": "광주 동구 동명로 1",
    "lat": 35.1495, "lng": 126.9285, "hours": "11:30 - 22:00", "closed": "월", "stay": 60, "cost": "9,000",
    "activity": "식사, 사진", "mood": "정겨운, 활기찬", "setting": "실내", "time": "오후, 저녁", "companions": "혼자, 친구",
    "aliases": "테스트국밥집", "desc": "2008년 개업.",
}
BAD = {**GOOD, "id": "s_902", "name": "잘못된 카페", "category": "술집", "hours": "열한시", "activity": "없는태그", "setting": "실내, 실외"}
NO_COORDS = {**GOOD, "id": "s_903", "name": "좌표 없는 가게", "lat": None, "lng": None}
OLD_MOOD = {**GOOD, "id": "s_904", "name": "옛 분위기", "mood": "서민적"}  # 새 분위기 목록에 없는 값


def _make(tmp_path, stores, checks=(), name="filled.xlsx"):
    path = tmp_path / name
    xlsx_template.build(str(path), stores)
    wb = openpyxl.load_workbook(path)
    sheet = wb["체크리스트"]
    marks = ["O", "x", "X", "o", "O", "O", "X", "O"]  # 수동 항목 8개 (history_site ... seating)
    for r, (sid, rater, m) in enumerate(checks, start=3):
        sheet[f"A{r}"], sheet[f"C{r}"] = sid, rater
        for c, v in zip("DEFGHIJK", m or marks):
            sheet[f"{c}{r}"] = v
        sheet[f"L{r}"] = "현장 확인"
    wb.save(path)
    return str(path)


@pytest.fixture
def filled(tmp_path):
    return _make(
        tmp_path,
        [GOOD, BAD, NO_COORDS, OLD_MOOD],
        [("s_901", "홍길동", None), ("s_901", "김철수", ["X"] * 8), ("s_999", "홍길동", None)],
    )


def test_dry_run_reports_each_problem_without_writing(filled):
    report = xlsx_import.load(filled, dry_run=True)
    assert report.places == ["s_901 테스트 국밥"]
    text = "\n".join(report.errors)
    assert "가게정보 4행 '잘못된 카페'" in text and "유형" in text and "술집" in text
    assert "가게정보 5행 '좌표 없는 가게'" in text and "좌표" in text
    assert "가게정보 6행 '옛 분위기'" in text and "서민적" in text
    assert "체크리스트 5행 s_999" in text
    with db.session() as s:
        assert s.get(db.PlaceRow, "s_901") is None  # dry-run 은 DB 를 바꾸지 않는다


def test_import_writes_places_checklists_and_scores(filled):
    try:
        report = xlsx_import.load(filled, placeholder=True)
        assert report.placeholders == ["s_903 좌표 없는 가게"]
        assert report.sheets == 2
        with db.session() as s:
            row = s.get(db.PlaceRow, "s_901")
            assert (row.name, row.category, row.open_time, row.close_time, row.closed_days) == ("테스트 국밥", "food", "11:30", "22:00", [0])
            assert (row.cost, row.stay_min, row.aliases) == (9000, 60, ["테스트국밥집"])
            assert row.tags["activity"] == ["사진", "식사"] and row.tags["setting"] == ["실내"]
            assert row.tags["mood"] == ["활기찬", "정겨운"]
            assert row.verified is False
            sheets = s.scalars(select(db.ChecklistRow).where(db.ChecklistRow.place_id == "s_901")).all()
            assert {c.rater for c in sheets} == {"홍길동", "김철수"}
            honggildong = next(c for c in sheets if c.rater == "홍길동")
            assert honggildong.checks["history_site"] is True and honggildong.checks["heritage"] is False
            assert honggildong.evidence["history_site"] == "현장 확인"
            # 광주다움 수동 4항목 중 홍길동이 O 인 것은 2개(history_site, long_running), 김철수는 전부 X -> 항목마다 0.5
            assert row.gwangju_ness == 1 + 0.5 + 0.5
            # not_franchise·market_zone 이 둘 다 0.5 (>= 0.5) 이므로 지역상권 장소
            assert row.is_local_commerce is True
        # 같은 파일을 다시 넣어도 중복되지 않는다 (덮어씀)
        xlsx_import.load(filled, placeholder=True)
        with db.session() as s:
            assert len(s.scalars(select(db.PlaceRow).where(db.PlaceRow.id == "s_901")).all()) == 1
            assert len(s.scalars(select(db.ChecklistRow).where(db.ChecklistRow.place_id == "s_901")).all()) == 2
    finally:
        _cleanup(["s_901", "s_903"])


def _cleanup(ids):
    with db.session() as s:
        s.execute(delete(db.ChecklistRow).where(db.ChecklistRow.place_id.in_(ids)))
        s.execute(delete(db.PlaceRow).where(db.PlaceRow.id.in_(ids)))
        s.commit()
    invalidate()


def test_replace_removes_candidates_but_keeps_anchors(tmp_path):
    path = _make(tmp_path, [GOOD], name="good.xlsx")
    with db.session() as s:
        n_candidates = len(s.scalars(select(db.PlaceRow).where(db.PlaceRow.kind == "candidate")).all())
        n_anchors = len(s.scalars(select(db.PlaceRow).where(db.PlaceRow.kind == "anchor")).all())
    assert n_candidates > 0 and n_anchors > 0
    try:
        # 고칠 곳이 있으면 아무것도 지우지 않는다
        r = xlsx_import.load(_make(tmp_path, [GOOD, BAD], name="bad.xlsx"), replace=True)
        assert r.errors
        with db.session() as s:
            assert len(s.scalars(select(db.PlaceRow).where(db.PlaceRow.kind == "candidate")).all()) == n_candidates

        r = xlsx_import.load(path, replace=True)
        assert not r.errors and r.removed == n_candidates
        with db.session() as s:
            ids = [p.id for p in s.scalars(select(db.PlaceRow).where(db.PlaceRow.kind == "candidate"))]
            assert ids == ["s_901"]
            assert len(s.scalars(select(db.PlaceRow).where(db.PlaceRow.kind == "anchor")).all()) == n_anchors
    finally:
        with db.session() as s:
            db.seed(s)  # 다른 테스트가 쓰는 예시 데이터를 되돌린다
        invalidate()


@pytest.mark.parametrize(
    "value, want",
    [("11:00", "11:00"), ("25:00", "25:00"), (" 9:30 ", "09:30"), (0.4583333333, "11:00"), (1.0416666667, "25:00")],
)
def test_time_cells(value, want):
    assert xlsx_import._hm(value) == want


@pytest.mark.parametrize(
    "value, want",
    [
        ("11:30-22:00", ("11:30", "22:00")),
        ("11:30 ~ 22:00", ("11:30", "22:00")),
        ("11:30 - 20:30", ("11:30", "20:30")),
        ("18:00-03:00", ("18:00", "27:00")),
        ("00:00-24:00", ("00:00", "24:00")),
        ("24시간", ("00:00", "24:00")),
        ("", (None, None)),
    ],
)
def test_hours_cell(value, want):
    assert xlsx_import._hours(value) == want


def test_hours_cell_rejects_garbage():
    for bad in ("열한시", "11:30", "11:30-22:00-23:00"):
        with pytest.raises(ValueError):
            xlsx_import._hours(bad)


def test_time_cell_from_excel_time_object():
    from datetime import time

    assert xlsx_import._hm(time(21, 5)) == "21:05"
    with pytest.raises(ValueError):
        xlsx_import._hm("열한시")


def test_closed_days_and_lists():
    assert xlsx_import._closed_days("월, 화요일") == [0, 1]
    assert xlsx_import._closed_days("연중무휴") == [] and xlsx_import._closed_days(None) == []
    assert xlsx_import._split("식사, 사진 / 산책\n야경") == ["식사", "사진", "산책", "야경"]
    assert xlsx_import._yes_no("O") is True and xlsx_import._yes_no("x") is False and xlsx_import._yes_no("") is None


def test_template_lists_follow_tags_json():
    from routefit import tags

    moods = tags.values("mood")
    assert moods == ["조용한", "활기찬", "캐주얼한", "정겨운", "아늑한", "고급스러운", "레트로한", "현대적인", "전통적인", "이국적인"]
    assert set(moods) == set(xlsx_template.MOOD_NOTES)  # 분위기 설명이 모든 값에 있다


def test_blank_template_dry_run_is_clean(tmp_path):
    path = tmp_path / "blank.xlsx"
    xlsx_template.build(str(path))
    report = xlsx_import.load(str(path), dry_run=True)
    assert report.places == [] and report.errors == []


def test_append_adds_to_empty_rows_skips_duplicates_and_keeps_dropdowns(tmp_path):
    path = str(tmp_path / "work.xlsx")
    xlsx_template.build(path, [GOOD])
    added, skipped = xlsx_template.append(path, [{"name": "테스트 국밥"}, {"name": "새 가게", "category": "식당"}])
    assert skipped == ["테스트 국밥"]
    assert added == [(4, "s_002", "새 가게")]  # s_001 줄 다음 빈 줄
    wb = openpyxl.load_workbook(path)
    assert wb["가게정보"]["B3"].value == "테스트 국밥"  # 기존 줄은 그대로
    assert len(wb["가게정보"].data_validations.dataValidation) == 2
    assert len(wb["체크리스트"].data_validations.dataValidation) == 2
