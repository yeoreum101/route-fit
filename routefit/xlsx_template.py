"""가게 조사 엑셀 양식을 만든다. 시트: 가이드 · 가게정보 · 체크리스트 · 목록.

    python -m routefit.xlsx_template docs/가게조사양식.xlsx

드롭다운과 '목록' 시트는 data/tags.json · data/checklist.json 에서 읽으므로, 목록을 바꾸면 이 명령으로 양식을 다시 만든다.
만든 엑셀은 `python -m routefit.xlsx_import` 로 DB 에 넣는다(열 머리글 이름으로 찾으므로 열 순서를 바꿔도 되지만 1행은 고치지 않는다).
"""
import sys
from typing import Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from . import checklist
from . import tags as tag_vocab

MAX_ROWS = 200  # 가게정보 데이터 줄 수 (3행부터)
MAX_SHEETS = 300  # 체크리스트 줄 수

CATEGORIES = ["식당", "카페·디저트", "전통시장", "상점가", "명소·산책", "문화·전시", "야간 콘텐츠"]
MOOD_NOTES = {
    "조용한": "차분하고 소음이 적은 분위기",
    "활기찬": "사람들의 대화와 에너지가 느껴지는 분위기",
    "캐주얼한": "격식 없이 편하게 이용할 수 있는 분위기",
    "정겨운": "친근하고 소박한 분위기",
    "아늑한": "포근하고 편안한 분위기",
    "고급스러운": "세련되고 격식 있는 분위기",
    "레트로한": "복고풍의 분위기",
    "현대적인": "현대적이고 세련된 디자인의 분위기",
    "전통적인": "전통문화와 고풍스러운 멋이 느껴지는 분위기",
    "이국적인": "다른 나라의 문화나 공간 특색이 느껴지는 분위기",
}

# (키, 머리글, 종류, 안내(2행), 너비)  종류: req 필수 · opt 선택 · ref 참고(저장 안 함) · auto 자동
COLUMNS = [
    ("id", "id\n(수정 금지)", "auto", "자동 번호", 11),
    ("name", "★가게 이름", "req", "간판 이름 그대로", 24),
    ("category", "★유형", "req", "목록에서 고르기", 13),
    ("area", "★지역", "req", "예: 동구 동명동", 14),
    ("address", "★도로명 주소", "req", "정확히. 좌표는 아래 위도·경도 또는 카카오 주소 검색으로", 34),
    ("lat", "위도", "opt", "몰라도 됨", 10),
    ("lng", "경도", "opt", "몰라도 됨", 10),
    ("hours", "★영업시간", "req", "예: 11:30-22:00 / 자정 넘기면 25:00 / 24시간은 00:00-24:00", 18),
    ("break", "브레이크타임\n[참고]", "ref", "예: 15:00-17:00", 14),
    ("last_order", "라스트 오더\n[참고]", "ref", "예: 20:30 / 14:30, 20:00", 14),
    ("closed", "휴무 요일", "opt", "월, 화 (연중무휴면 비움)", 12),
    ("closed_note", "휴무 예외\n[참고]", "ref", "예: 재료 소진 시 조기 마감", 24),
    ("stay", "★머무는 시간(분)", "req", "보통 머무는 시간 예: 60", 12),
    ("cost", "★1인 예상 지출(원)", "req", "메뉴 1개+음료 기준. 무료는 0", 13),
    ("food_kind", "음식 종류\n[참고]", "ref", "식당·카페만", 14),
    ("menu", "대표 메뉴·가격\n[참고]", "ref", "예: 순대국밥 9,000 / 수육 25,000", 34),
    ("features", "편의·특징\n[참고]", "ref", "예: 주차 가능, 예약 가능", 24),
    ("activity", "★할 거리", "req", "목록 값을 쉼표로 예: 식사, 사진", 20),
    ("mood", "★분위기", "req", "목록 값을 쉼표로 (2~3개 권장)", 26),
    ("setting", "★공간", "req", "하나만 고르기", 10),
    ("time", "어울리는 시간", "opt", "쉼표로. 예: 저녁, 밤", 16),
    ("companions", "어울리는 동행", "opt", "쉼표로. 예: 친구, 연인", 18),
    ("aliases", "별칭", "opt", "손님들이 부르는 다른 이름 (쉼표)", 18),
    ("desc", "한줄 소개", "opt", "확인된 사실만. 홍보 문구·추측 금지", 40),
    ("link", "네이버 링크", "ref", "https://map.naver.com/p/entry/place/…", 30),
    ("checked", "확인한 날짜", "ref", "예: 2026-10-21", 12),
    ("rater", "조사자", "ref", "이름", 10),
    ("memo", "메모", "ref", "", 30),
]
KEYS = [c[0] for c in COLUMNS]

GREEN = PatternFill("solid", fgColor="1E6B4E")
SLATE = PatternFill("solid", fgColor="5B6B7B")
LIGHT = PatternFill("solid", fgColor="BFD8F2")
GRAY = PatternFill("solid", fgColor="E7E7E7")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEADER_FILL = {"req": GREEN, "opt": SLATE, "ref": LIGHT, "auto": GRAY}


def _header(cell, kind: str) -> None:
    cell.fill = HEADER_FILL[kind]
    cell.font = Font(bold=True, color="FFFFFF" if kind in ("req", "opt") else "222222")
    cell.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
    cell.border = BORDER


def _hint(cell) -> None:
    cell.fill = GRAY
    cell.font = Font(italic=True, size=9, color="555555")
    cell.alignment = Alignment(wrap_text=True, vertical="top")
    cell.border = BORDER


def _list(ws, ref: str, formula: str, title: str, error: str, strict: bool = True) -> None:
    dv = DataValidation(type="list", formula1=formula, allow_blank=True, showErrorMessage=True, errorTitle=title, error=error)
    dv.errorStyle = "stop" if strict else "warning"
    ws.add_data_validation(dv)
    dv.add(ref)


def _norm(v) -> str:
    return "".join(str(v or "").replace("★", "").split())


def _columns(ws) -> dict[str, int]:
    """가게정보 시트의 머리글 -> 키 -> 열 번호(1부터). 열 순서를 바꿨어도 머리글 이름으로 찾는다."""
    want = {_norm(title): key for key, title, *_ in COLUMNS}
    out = {}
    for c in range(1, ws.max_column + 1):
        key = want.get(_norm(ws.cell(1, c).value))
        if key:
            out[key] = c
    return out


def add_dropdowns(wb) -> None:
    """유형·공간·체크리스트 O/X·가게 id 드롭다운을 (다시) 건다.
    openpyxl 은 엑셀이 다시 저장한 파일의 '다른 시트 참조' 드롭다운을 읽지 못하고 지워 버리므로, 파일을 고쳐 저장한 뒤에는 이걸 다시 부른다."""
    info, sheet = wb["가게정보"], wb["체크리스트"]
    for ws in (info, sheet):
        for dv in list(ws.data_validations.dataValidation):
            ws.data_validations.dataValidation.remove(dv)
    cols = _columns(info)
    last = 2 + MAX_ROWS
    n_setting = len(tag_vocab.groups()[[g["key"] for g in tag_vocab.groups()].index("setting")]["values"])
    letter = get_column_letter
    _list(info, f"{letter(cols['category'])}3:{letter(cols['category'])}{last}", f"=목록!$A$2:$A${1 + len(CATEGORIES)}", "유형", "목록에서 골라 주세요.")
    _list(info, f"{letter(cols['setting'])}3:{letter(cols['setting'])}{last}", f"=목록!$B$2:$B${1 + n_setting}", "공간", "실내 / 실외 / 실내외 중 하나만 고르세요.")
    n_manual = len([it for it in checklist.definition()["items"] if not it.get("auto")])
    _list(sheet, f"A3:A{2 + MAX_SHEETS}", f"=가게정보!$A$3:$A${last}", "가게 id", "가게정보의 id 중에서 골라 주세요.")
    _list(sheet, f"D3:{letter(3 + n_manual)}{2 + MAX_SHEETS}", "=목록!$C$2:$C$3", "O/X", "O 또는 X 로 적어 주세요.")


def append(path: str, stores: list[dict]) -> tuple[list[tuple[int, str, str]], list[str]]:
    """이미 채우고 있는 엑셀의 빈 줄에 가게를 더한다. 같은 이름(공백 무시)이 있으면 건너뛴다.
    돌려주는 값: ([(줄, id, 이름)], [건너뛴 이름]). 기존 칸은 건드리지 않는다."""
    from openpyxl import load_workbook

    wb = load_workbook(path)
    ws = wb["가게정보"]
    cols = _columns(ws)
    names = {_norm(ws.cell(r, cols["name"]).value) for r in range(3, ws.max_row + 1) if ws.cell(r, cols["name"]).value}
    added, skipped = [], []
    row = 3
    for rec in stores:
        if _norm(rec["name"]) in names:
            skipped.append(rec["name"])
            continue
        while ws.cell(row, cols["name"]).value:
            row += 1
        for key, value in rec.items():
            if key in cols and value not in (None, ""):
                ws.cell(row, cols[key], value)
        added.append((row, str(ws.cell(row, cols["id"]).value), rec["name"]))
        names.add(_norm(rec["name"]))
    add_dropdowns(wb)
    wb.save(path)
    return added, skipped


def build(path: str, stores: Optional[list[dict]] = None) -> None:
    """stores: [{키: 값}] 을 3행부터 미리 채운다(키는 COLUMNS 의 첫 칸). 없으면 빈 양식."""
    wb = Workbook()
    guide = wb.active
    guide.title = "가이드"
    info = wb.create_sheet("가게정보")
    sheet = wb.create_sheet("체크리스트")
    lists = wb.create_sheet("목록")

    # ----- 목록 -----
    groups = {g["key"]: g for g in tag_vocab.groups()}
    list_cols = [
        ("유형", CATEGORIES),
        ("공간", groups["setting"]["values"]),
        ("O/X", ["O", "X"]),
        ("할 거리 (여러 개 가능)", groups["activity"]["values"]),
        ("분위기 (여러 개 가능)", groups["mood"]["values"]),
        ("분위기 설명", [MOOD_NOTES.get(v, "") for v in groups["mood"]["values"]]),
        ("어울리는 시간 (여러 개 가능)", groups["time"]["values"]),
        ("어울리는 동행 (여러 개 가능)", groups["companions"]["values"]),
    ]
    for c, (title, values) in enumerate(list_cols, 1):
        _header(lists.cell(1, c, title), "opt")
        lists.column_dimensions[get_column_letter(c)].width = 40 if title == "분위기 설명" else 22
        for r, v in enumerate(values, 2):
            lists.cell(r, c, v)

    # ----- 가게정보 -----
    info.row_dimensions[1].height = 42
    info.row_dimensions[2].height = 52
    for c, (key, title, kind, hint, width) in enumerate(COLUMNS, 1):
        _header(info.cell(1, c, title), kind)
        _hint(info.cell(2, c, hint))
        info.column_dimensions[get_column_letter(c)].width = width
    text_cols = {KEYS.index(k) + 1 for k in ("hours", "break", "last_order", "closed", "checked")}
    for i in range(MAX_ROWS):
        row = 3 + i
        for c in range(1, len(COLUMNS) + 1):
            cell = info.cell(row, c)
            cell.border = BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if c in text_cols:
                cell.number_format = "@"  # 엑셀이 11:30 을 시각 값으로 바꾸지 않게
        id_cell = info.cell(row, 1, f"s_{i + 1:03d}")
        id_cell.fill = GRAY
    info.freeze_panes = "C3"
    last = 2 + MAX_ROWS

    for i, store in enumerate(stores or []):
        for key, value in store.items():
            if key in KEYS and value not in (None, ""):
                info.cell(3 + i, KEYS.index(key) + 1, value)

    # ----- 체크리스트 -----
    manual = [it for it in checklist.definition()["items"] if not it.get("auto")]
    attr_label = {a["key"]: a["label"] for a in checklist.definition()["attributes"]}
    heads = [("가게 id", "auto"), ("가게 이름", "auto"), ("★채점자", "req")]
    heads += [(f"[{attr_label[it['attribute']]}]\n{it['label']}", "req") for it in manual]
    heads += [("근거\n(링크·사진 파일명·메모)", "opt")]
    hints = ["가게정보의 id 고르기", "자동", "본인 이름"] + [f"확인 방법: {it['how']}" for it in manual] + ["O/X를 정한 근거. 항목 이름과 함께 적기"]
    sheet.row_dimensions[1].height = 78
    sheet.row_dimensions[2].height = 64
    for c, ((title, kind), hint) in enumerate(zip(heads, hints), 1):
        _header(sheet.cell(1, c, title), kind)
        _hint(sheet.cell(2, c, hint))
        sheet.column_dimensions[get_column_letter(c)].width = 14 if c > 2 and c < len(heads) else (22 if c <= 2 else 36)
    for r in range(3, 3 + MAX_SHEETS):
        sheet.cell(r, 2, f'=IF(A{r}="","",IFERROR(INDEX(가게정보!$B$3:$B${last},MATCH(A{r},가게정보!$A$3:$A${last},0)),"(id를 확인하세요)"))').fill = GRAY
        for c in range(1, len(heads) + 1):
            sheet.cell(r, c).border = BORDER
    sheet.freeze_panes = "D3"
    add_dropdowns(wb)

    # ----- 가이드 -----
    rows = [
        ("Route-Fit 가게 조사 양식", None),
        ("가게 한 곳 = 가게정보 시트의 한 줄입니다. 네이버 지도 등에서 확인한 값을 적고, 확인 못 한 값은 비워 두세요(추측 금지).", None),
        ("1. 시트", None),
        ("가게정보", "한 줄 = 가게 1곳. 2행 안내는 지우지 마세요. id(A열)는 고치지 마세요."),
        ("체크리스트", "광주다움·상권 연계성·체류 유도성의 수동 항목 8개를 O/X 로 체크. 가게마다 팀원 2명 이상이 서로 보지 않고 각자 한 줄씩."),
        ("목록", "드롭다운과 태그 값 목록. 할 거리·분위기·시간·동행은 이 값만 쉼표로 적습니다."),
        ("2. 칸 색", None),
        ("진한 초록 머리글 (★)", "필수. 비우면 DB 에 넣을 수 없습니다."),
        ("회청색 머리글", "선택. 아는 만큼만."),
        ("연한 파랑 머리글 [참고]", "브레이크타임·라스트 오더·메뉴·편의·링크·확인 날짜 등. 지금 DB 에는 저장하지 않지만 확인용으로 남겨 둡니다."),
        ("3. 작성 규칙", None),
        ("영업시간", "11:30-22:00 처럼 24시간제. 자정을 넘기면 25:00(새벽 1시). 24시간 영업은 00:00-24:00. 요일마다 다르면 가장 흔한 시간을 적고 메모에 예외를 남기세요."),
        ("휴무 요일", "월, 화, 수, 목, 금, 토, 일 중 쉼표로. 연중무휴면 비움. '재료 소진 시 조기 마감' 같은 예외는 휴무 예외에."),
        ("1인 예상 지출", "메뉴 1개 + 음료 기준 대략의 값(원). 무료는 0."),
        ("할 거리 · 공간", "필수입니다. 목록 시트의 값만 씁니다(식당이면 보통 '식사', 공간은 실내/실외/실내외 중 하나)."),
        ("분위기", "목록 시트의 10개 중에서 2~3개를 고르세요. 많이 고를수록 모든 가게가 비슷해져 취향 매칭이 흐려집니다. 분위기 설명은 목록 시트에 있습니다."),
        ("한줄 소개", "확인된 사실만(대표 메뉴, 개업 연도 등). 홍보 문구·추측 금지."),
        ("주소·좌표", "도로명 주소만 적어도 됩니다. 위도·경도는 비워 두고 가져올 때 --geocode(카카오 키 필요)나 장소 관리 화면의 지도 클릭으로 넣습니다."),
        ("4. 체크리스트", None),
        ("O / X", "'맞다/아니다'를 사실로 확인합니다. 확인 방법은 체크리스트 시트 2행에 있습니다. 근거(링크·사진 파일명)를 적어 주세요."),
        ("자동 항목", "머무는 시간 60분 이상, 지출 0원 초과, 할 거리 2개 이상, 19시 이후 운영은 시스템이 계산하므로 적지 않습니다."),
        ("프랜차이즈", "'프랜차이즈나 대기업 직영이 아니다'는 프랜차이즈가 아니면 O 입니다."),
        ("5. DB 에 넣기", None),
        ("검사", "python -m routefit.xlsx_import 파일.xlsx --dry-run  (틀린 줄은 몇 행 어느 칸인지 알려 줍니다)"),
        ("넣기", "python -m routefit.xlsx_import 파일.xlsx  (예시 데이터를 지우고 이 파일로 바꾸려면 --replace)"),
    ]
    guide.column_dimensions["A"].width = 26
    guide.column_dimensions["B"].width = 110
    for r, (a, b) in enumerate(rows, 1):
        ca, cb = guide.cell(r, 1, a), guide.cell(r, 2, b)
        cb.alignment = Alignment(wrap_text=True, vertical="top")
        ca.alignment = Alignment(wrap_text=True, vertical="top")
        if b is None:
            ca.font = Font(bold=True, size=14 if r == 1 else 12)
        else:
            ca.font = Font(bold=True)
    guide.merge_cells("A1:B1")
    guide.merge_cells("A2:B2")
    guide["A2"].alignment = Alignment(wrap_text=True, vertical="top")
    guide.row_dimensions[2].height = 32

    wb.save(path)


if __name__ == "__main__":
    out = sys.argv[1] if len(sys.argv) > 1 else "가게조사양식.xlsx"
    build(out)
    print(f"만들었습니다: {out}")
